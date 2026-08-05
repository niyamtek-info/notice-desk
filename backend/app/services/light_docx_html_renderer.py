from __future__ import annotations

import base64
import mimetypes
import os
import re
import zipfile
from dataclasses import dataclass
from html import escape
from typing import Optional

from docx import Document
from lxml import etree

from app.core.settings import settings


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"w": W_NS, "r": R_NS, "a": A_NS}
PLACEHOLDER_RE = re.compile(r"<\s*([^<>]+?)\s*>")


def _local_name(elem) -> str:
    return etree.QName(elem).localname


def _twips_to_px(value: Optional[str]) -> int:
    if not value:
        return 0
    try:
        return max(int(value) // 20, 0)
    except ValueError:
        return 0


def _length_to_px(value) -> int:
    if value is None:
        return 0


def _emu_to_px(value: Optional[str]) -> int:
    if not value:
        return 0
    try:
        return max(int(round(int(value) / 9525)), 0)
    except Exception:
        return 0
    try:
        if hasattr(value, "pt") and value.pt is not None:
            return max(int(round(float(value.pt) * 1.333333)), 0)
    except Exception:
        pass
    try:
        return max(int(value), 0)
    except Exception:
        return 0


def _normalize_placeholder_name(raw: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", raw.upper()).strip("_")
    return re.sub(r"_+", "_", normalized)


@dataclass
class TableMeta:
    visible_borders: bool
    col_widths: list[int]


class LightDocxHtmlRenderer:
    def __init__(self, file_path: str):
        self.file_path = self._resolve_path(file_path)
        self.doc = Document(self.file_path)
        self.document_xml, self.part_rels, self.media = self._load_package(self.file_path)

    def _resolve_path(self, file_path: str) -> str:
        candidates = [
            file_path,
            os.path.join(os.getcwd(), file_path),
            os.path.join(settings.BASE_DIR, file_path),
            os.path.join(settings.BASE_DIR, "app", file_path),
        ]
        for candidate in candidates:
            if candidate and os.path.exists(candidate):
                return os.path.abspath(candidate)
        return os.path.abspath(file_path)

    def _load_package(self, file_path: str):
        part_rels: dict[str, dict[str, dict[str, str]]] = {}
        media: dict[str, bytes] = {}
        with zipfile.ZipFile(file_path) as zf:
            document_xml = etree.fromstring(zf.read("word/document.xml"))
            for name in zf.namelist():
                if name.startswith("word/_rels/") and name.endswith(".rels"):
                    rels_xml = etree.fromstring(zf.read(name))
                    owner = f"word/{name[len('word/_rels/'):-len('.rels')]}"
                    rels: dict[str, dict[str, str]] = {}
                    for rel in rels_xml.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                        rel_id = rel.attrib.get("Id")
                        if not rel_id:
                            continue
                        rels[rel_id] = {
                            "type": rel.attrib.get("Type", ""),
                            "target": rel.attrib.get("Target", ""),
                            "mode": rel.attrib.get("TargetMode", ""),
                        }
                    part_rels[owner] = rels
            for name in zf.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    media[name.split("word/", 1)[1]] = zf.read(name)
        return document_xml, part_rels, media

    def _part_rels(self, part) -> dict[str, dict[str, str]]:
        part_name = str(getattr(part, "partname", "") or "").lstrip("/")
        return self.part_rels.get(part_name, {})

    def _escape_with_placeholders(self, value: str) -> str:
        if not value:
            return ""
        pieces: list[str] = []
        index = 0
        for match in PLACEHOLDER_RE.finditer(value):
            pieces.append(escape(value[index:match.start()]))
            pieces.append(f"{{{{{_normalize_placeholder_name(match.group(1))}}}}}")
            index = match.end()
        pieces.append(escape(value[index:]))
        return "".join(pieces)

    def _normalize_mojibake(self, value: str) -> str:
        if not value:
            return value
        replacements = {
            "â€œ": "“",
            "â€": "”",
            "â€˜": "‘",
            "â€™": "’",
            "â€“": "–",
            "â€”": "—",
            "â€¦": "…",
            "Â": "",
        }
        for source, target in replacements.items():
            value = value.replace(source, target)
        return value

    def _repair_split_placeholders(self, html_text: str) -> str:
        if not html_text or "{{" not in html_text or "}}" not in html_text:
            return html_text

        def repl(match: re.Match) -> str:
            inner = re.sub(r"<[^>]+>", "", match.group(0))
            inner = inner.replace("{", "").replace("}", "")
            inner = re.sub(r"\s+", " ", inner).strip()
            if not inner:
                return match.group(0)
            return f"{{{{{_normalize_placeholder_name(inner)}}}}}"

        return re.sub(r"\{\{.*?\}\}", repl, html_text, flags=re.S)

    def _resolve_media_html(
        self,
        rel_id: str,
        rels: dict[str, dict[str, str]],
        width_px: int = 0,
        height_px: int = 0,
    ) -> str:
        rel = rels.get(rel_id)
        if not rel:
            return ""
        target = rel.get("target", "").lstrip("/")
        if not target:
            return ""
        media_name = target if target.startswith("word/") else f"word/{target}"
        media_name = media_name.split("word/", 1)[1]
        data = self.media.get(media_name)
        if not data:
            return ""
        mime = mimetypes.guess_type(media_name)[0] or "application/octet-stream"
        style_bits = []
        if width_px > 0:
            style_bits.append(f"width:{width_px}px")
        if height_px > 0:
            style_bits.append(f"height:{height_px}px")
        style_attr = f' style="{"; ".join(style_bits)}"' if style_bits else ""
        return f'<img src="data:{mime};base64,{base64.b64encode(data).decode("ascii")}"{style_attr} />'

    def _run_html(self, run_elem, rels: dict[str, dict[str, str]]) -> str:
        pieces: list[str] = []
        text_parts: list[str] = []
        image_html = ""

        def flush() -> None:
            nonlocal text_parts, image_html
            text = self._normalize_mojibake("".join(text_parts))
            if text or image_html:
                rendered = self._escape_with_placeholders(text).replace("\n", "<br/>")
                if run_elem.find("./w:rPr/w:u", NS) is not None:
                    rendered = f"<u>{rendered}</u>"
                if run_elem.find("./w:rPr/w:i", NS) is not None:
                    rendered = f"<em>{rendered}</em>"
                if run_elem.find("./w:rPr/w:b", NS) is not None:
                    rendered = f"<strong>{rendered}</strong>"
                if image_html:
                    rendered += image_html
                pieces.append(rendered)
            text_parts = []
            image_html = ""

        for child in list(run_elem):
            tag = _local_name(child)
            if tag == "t":
                text_parts.append(child.text or "")
            elif tag == "tab":
                text_parts.append(" ")
            elif tag in {"br", "cr"}:
                text_parts.append("\n")
            elif tag == "drawing":
                blip = child.find(".//a:blip", NS)
                if blip is not None:
                    rel_id = blip.attrib.get(f"{{{R_NS}}}embed")
                    if rel_id:
                        extent = None
                        for descendant in child.iter():
                            if _local_name(descendant) == "extent":
                                cx = descendant.attrib.get("cx")
                                cy = descendant.attrib.get("cy")
                                if cx or cy:
                                    extent = descendant
                                    break
                        width_px = _emu_to_px(extent.attrib.get("cx")) if extent is not None else 0
                        height_px = _emu_to_px(extent.attrib.get("cy")) if extent is not None else 0
                        image_html = self._resolve_media_html(rel_id, rels, width_px=width_px, height_px=height_px)

        flush()
        return "".join(pieces)

    def _paragraph_css(self, p_elem, compact: bool = False) -> str:
        ppr = p_elem.find("./w:pPr", NS)
        styles = ["margin:0" if compact else "margin:0 0 6px 0", "white-space:pre-wrap"]
        if ppr is None:
            return "; ".join(styles)

        jc = ppr.find("./w:jc", NS)
        if jc is not None:
            align = jc.attrib.get(f"{{{W_NS}}}val") or jc.attrib.get("val")
            align = align if align in {"left", "center", "right", "justify", "both"} else "left"
            if align == "both":
                align = "left" if compact else "justify"
            styles.append(f"text-align:{align}")

        spacing = ppr.find("./w:spacing", NS)
        if spacing is not None:
            before = spacing.attrib.get(f"{{{W_NS}}}before")
            after = spacing.attrib.get(f"{{{W_NS}}}after")
            line = spacing.attrib.get(f"{{{W_NS}}}line")
            if before:
                styles.append(f"margin-top:{_twips_to_px(before)}px")
            if after:
                styles.append(f"margin-bottom:{_twips_to_px(after)}px")
            if line:
                try:
                    styles.append(f"line-height:{max(int(line) / 240.0, 1.0):.2f}")
                except ValueError:
                    pass

        ind = ppr.find("./w:ind", NS)
        if ind is not None:
            left = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}left"))
            first = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}firstLine"))
            hanging = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}hanging"))
            if left:
                styles.append(f"margin-left:{left}px")
            if first:
                styles.append(f"text-indent:{first}px")
            if hanging:
                styles.append(f"text-indent:-{hanging}px")
                if left:
                    styles.append(f"margin-left:{left + hanging}px")

        return "; ".join(styles)

    def _paragraph_html(self, p_elem, rels: dict[str, dict[str, str]], compact: bool = False) -> str:
        content = []
        for child in list(p_elem):
            tag = _local_name(child)
            if tag == "r":
                rendered = self._run_html(child, rels)
                if rendered:
                    content.append(rendered)
            elif tag == "hyperlink":
                for run in child.findall("./w:r", NS):
                    rendered = self._run_html(run, rels)
                    if rendered:
                        content.append(rendered)
            elif tag == "smartTag":
                for run in child.findall(".//w:r", NS):
                    rendered = self._run_html(run, rels)
                    if rendered:
                        content.append(rendered)
        html = "".join(content)
        if not html.strip():
            return ""
        visible_text = re.sub(r"<[^>]+>", "", html).strip()
        if not visible_text and '<img ' not in html:
            return ""
        html = self._repair_split_placeholders(html)
        return f'<p style="{self._paragraph_css(p_elem, compact=compact)}">{html}</p>'

    def _table_has_visible_borders(self, table_elem) -> bool:
        tbl_pr = table_elem.find("./w:tblPr", NS)
        if tbl_pr is None:
            return False
        borders = tbl_pr.find("./w:tblBorders", NS)
        if borders is not None:
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = borders.find(f"./w:{side}", NS)
                if border is None:
                    continue
                value = (border.attrib.get(f"{{{W_NS}}}val") or border.attrib.get("val") or "").lower()
                if value and value not in {"nil", "none", "0"}:
                    return True

        for tc in table_elem.findall(".//w:tc", NS):
            tc_pr = tc.find("./w:tcPr", NS)
            if tc_pr is None:
                continue
            tc_borders = tc_pr.find("./w:tcBorders", NS)
            if tc_borders is None:
                continue
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = tc_borders.find(f"./w:{side}", NS)
                if border is None:
                    continue
                value = (border.attrib.get(f"{{{W_NS}}}val") or border.attrib.get("val") or "").lower()
                if value and value not in {"nil", "none", "0"}:
                    return True
        return False

    def _table_meta(self, table_elem) -> TableMeta:
        visible_borders = self._table_has_visible_borders(table_elem)
        col_widths: list[int] = []
        tbl_grid = table_elem.find("./w:tblGrid", NS)
        if tbl_grid is not None:
            for col in tbl_grid.findall("./w:gridCol", NS):
                width = col.attrib.get(f"{{{W_NS}}}w")
                if width:
                    try:
                        col_widths.append(max(int(width) // 20, 1))
                    except ValueError:
                        col_widths.append(0)
        return TableMeta(visible_borders=visible_borders, col_widths=col_widths)

    def _cell_html(self, cell_elem, rels: dict[str, dict[str, str]]) -> str:
        blocks: list[str] = []
        for child in list(cell_elem):
            tag = _local_name(child)
            if tag == "p":
                rendered = self._paragraph_html(child, rels, compact=True)
                if rendered:
                    blocks.append(rendered)
            elif tag == "tbl":
                blocks.append(self._table_html(child, rels))
        return "".join(blocks) or "&nbsp;"

    def _table_html(self, table_elem, rels: dict[str, dict[str, str]]) -> str:
        meta = self._table_meta(table_elem)
        rows_html: list[str] = []
        tr_elems = [child for child in list(table_elem) if _local_name(child) == "tr"]

        for row_idx, row in enumerate(tr_elems):
            tc_elems = [child for child in list(row) if _local_name(child) == "tc"]
            cells: list[str] = []
            for cell_idx, cell in enumerate(tc_elems):
                tc_pr = cell.find("./w:tcPr", NS)
                colspan = 1
                rowspan = None
                if tc_pr is not None:
                    grid_span = tc_pr.find("./w:gridSpan", NS)
                    if grid_span is not None:
                        try:
                            colspan = max(int(grid_span.attrib.get(f"{{{W_NS}}}val", "1")), 1)
                        except ValueError:
                            colspan = 1
                    vmerge = tc_pr.find("./w:vMerge", NS)
                    if vmerge is not None:
                        vmerge_val = vmerge.attrib.get(f"{{{W_NS}}}val")
                        if vmerge_val == "restart" or vmerge_val is None:
                            rowspan = 1
                        else:
                            continue

                cell_html = self._cell_html(cell, rels)
                tag = "th" if meta.visible_borders and row_idx == 0 else "td"
                align = "center" if meta.visible_borders and row_idx == 0 else "left"
                if tc_pr is not None:
                    jc = tc_pr.find("./w:jc", NS)
                    if jc is not None:
                        align = jc.attrib.get(f"{{{W_NS}}}val") or align
                if not meta.visible_borders:
                    align = "left"
                border_css = "border:1px solid #111827;" if meta.visible_borders else "border:none;"
                width_css = ""
                if cell_idx < len(meta.col_widths) and meta.col_widths[cell_idx]:
                    width_css = f" width:{meta.col_widths[cell_idx]}px;"
                attrs = [
                    f'style="{border_css} padding:8px 10px; vertical-align:top; text-align:{align}; line-height:1.2;{width_css}"'
                ]
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')
                if rowspan:
                    attrs.append(f'rowspan="{rowspan}"')
                cells.append(f"<{tag} {' '.join(attrs)}>{cell_html}</{tag}>")
            rows_html.append(f"<tr>{''.join(cells)}</tr>")

        table_class = "docx-table" if meta.visible_borders else "docx-table hidden-table"
        return (
            f'<table class="{table_class}" style="width:100%; border-collapse:collapse; '
            f'table-layout:fixed; margin:0; page-break-inside:avoid;">{"".join(rows_html)}</table>'
        )

    def _length_to_mm(self, value, default_mm: float) -> float:
        if value is None:
            return default_mm
        try:
            if hasattr(value, "mm") and value.mm is not None:
                return float(value.mm)
            if hasattr(value, "pt") and value.pt is not None:
                return float(value.pt) * 0.352778
        except Exception:
            pass
        return default_mm

    def _render_container(self, container, rels: dict[str, dict[str, str]]) -> str:
        root = getattr(container, "_element", None)
        if root is None:
            return ""
        parts: list[str] = []
        for child in list(root):
            tag = _local_name(child)
            if tag == "p":
                rendered = self._paragraph_html(child, rels, compact=False)
                if rendered:
                    parts.append(rendered)
            elif tag == "tbl":
                parts.append(self._table_html(child, rels))
        return "\n".join(parts)

    def _iter_header_footers(self):
        seen = set()
        for section in self.doc.sections:
            for container in (
                getattr(section, "header", None),
                getattr(section, "first_page_header", None),
                getattr(section, "even_page_header", None),
                getattr(section, "footer", None),
                getattr(section, "first_page_footer", None),
                getattr(section, "even_page_footer", None),
            ):
                if container is None:
                    continue
                key = id(getattr(container, "_element", container))
                if key in seen:
                    continue
                seen.add(key)
                yield container

    def _render_header_footer_blocks(self) -> tuple[str, str]:
        header_blocks: list[str] = []
        footer_blocks: list[str] = []
        seen = set()
        for section in self.doc.sections:
            for container, is_header in (
                (getattr(section, "header", None), True),
                (getattr(section, "first_page_header", None), True),
                (getattr(section, "even_page_header", None), True),
                (getattr(section, "footer", None), False),
                (getattr(section, "first_page_footer", None), False),
                (getattr(section, "even_page_footer", None), False),
            ):
                if container is None:
                    continue
                key = id(getattr(container, "_element", container))
                if key in seen:
                    continue
                seen.add(key)
                rels = self._part_rels(getattr(container, "part", None))
                rendered = self._render_container(container, rels)
                if rendered.strip():
                    if is_header:
                        header_blocks.append(f'<header class="docx-header">{rendered}</header>')
                    else:
                        footer_blocks.append(f'<footer class="docx-footer">{rendered}</footer>')
        return "\n".join(header_blocks), "\n".join(footer_blocks)

    def render(self) -> str:
        body = self.document_xml.find("./w:body", NS)
        if body is None:
            return ""
        parts: list[str] = []
        children = list(body)
        for idx, child in enumerate(children):
            tag = _local_name(child)
            if tag == "p":
                rendered = self._paragraph_html(child, self._part_rels(self.doc.part))
                if rendered:
                    parts.append(rendered)
            elif tag == "tbl":
                parts.append(self._table_html(child, self._part_rels(self.doc.part)))
            elif tag == "sectPr":
                if idx < len(children) - 1:
                    parts.append('<div style="break-after:page;"></div>')
        return "\n".join(parts)


def docx_to_html(file_path: str) -> str:
    renderer = LightDocxHtmlRenderer(file_path)
    body_html = renderer.render()
    header_html, footer_html = renderer._render_header_footer_blocks()
    first_section = renderer.doc.sections[0] if renderer.doc.sections else None
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <style>
    * {{
      box-sizing: border-box;
    }}

    body {{
      margin: 0;
      font-family: Arial, Helvetica, sans-serif;
      font-size: 12pt;
      line-height: 1.5;
      color: #111827;
      background: #fff;
    }}

    .docx-root {{
      width: 100%;
    }}

    .docx-header,
    .docx-content,
    .docx-footer {{
      width: 100%;
      display: block;
    }}

    .docx-table {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
    }}

    .docx-table.hidden-table,
    .docx-table.hidden-table th,
    .docx-table.hidden-table td {{
      border: none;
    }}

    .docx-table th,
    .docx-table td {{
      vertical-align: top;
    }}

    p, li, h1, h2, h3, h4, h5, h6 {{
      margin: 0;
      white-space: pre-wrap;
    }}

    img {{
      max-width: 100%;
      height: auto;
    }}

    .docx-header img {{
      display: block;
      width: auto !important;
      height: auto !important;
      max-height: 25mm !important;
      max-width: 100% !important;
      object-fit: contain;
      margin: 0 auto;
    }}

    .docx-footer img {{
      display: block;
      width: auto !important;
      height: auto !important;
      max-height: 25mm !important;
      max-width: 100% !important;
      object-fit: contain;
      margin: 0 auto !important;
    }}

    .docx-header table,
    .docx-footer table {{
      width: 100% !important;
      margin: 0;
      border-collapse: collapse;
      table-layout: fixed;
    }}

    .docx-header p,
    .docx-footer p {{
      margin: 0 !important;
    }}

    .docx-footer p {{
      text-align: center !important;
    }}

  </style>
</head>
<body>
  <div class="docx-root">
    {header_html}
    <div class="docx-content">
      {body_html}
    </div>
    {footer_html}
  </div>
</body>
</html>"""
