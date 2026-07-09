from __future__ import annotations

import base64
import mimetypes
import os
import re
import zipfile
from dataclasses import dataclass
from html import escape
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

from app.core.settings import settings


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {
    "w": W_NS,
    "r": R_NS,
    "a": A_NS,
    "wp": WP_NS,
}

PLACEHOLDER_RE = re.compile(r"<\s*([^<>]+?)\s*>")


def _local_name(elem) -> str:
    return etree.QName(elem).localname


def _twips_to_px(value: Optional[str]) -> int:
    if not value:
        return 0
    try:
        return max(int(value) // 20, 0)
    except ValueError:
        return 0


def _half_points_to_pt(value: Optional[str]) -> Optional[float]:
    if not value:
        return None
    try:
        return int(value) / 2.0
    except ValueError:
        return None


def _normalize_placeholder_name(raw: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", raw.upper()).strip("_")
    return re.sub(r"_+", "_", normalized)


@dataclass
class ListFormat:
    tag: str
    type_attr: str = ""


class DocxHtmlRenderer:
    def __init__(self, file_path: str):
        self.file_path = self._resolve_path(file_path)
        self._doc = Document(self.file_path)
        self.document_xml, self.styles_xml, self.numbering_xml, self.part_rels, self.media = self._load_package(
            self.file_path
        )
        self.styles = self._parse_styles()
        self.numbering = self._parse_numbering()
        self.rels = self.part_rels.get("word/document.xml", {})

    def _resolve_path(self, file_path: str) -> str:
        candidates = [
            file_path,
            os.path.join(os.getcwd(), file_path),
            os.path.join(settings.BASE_DIR, file_path),
            os.path.join(settings.BASE_DIR, "app", file_path),
        ]
        for candidate in candidates:
            if candidate and os.path.exists(candidate):
                return os.path.abspath(candidate)
        return os.path.abspath(file_path)

    def _load_package(self, file_path: str):
        styles_xml = None
        numbering_xml = None
        part_rels: dict[str, dict[str, dict[str, str]]] = {}
        media: dict[str, bytes] = {}

        with zipfile.ZipFile(file_path) as zf:
            document_xml = etree.fromstring(zf.read("word/document.xml"))
            if "word/styles.xml" in zf.namelist():
                styles_xml = etree.fromstring(zf.read("word/styles.xml"))
            if "word/numbering.xml" in zf.namelist():
                numbering_xml = etree.fromstring(zf.read("word/numbering.xml"))
            for name in zf.namelist():
                if name.startswith("word/_rels/") and name.endswith(".rels"):
                    rels_xml = etree.fromstring(zf.read(name))
                    owner = name[len("word/_rels/"):-len(".rels")]
                    owner = f"word/{owner}"
                    rels = {}
                    for rel in rels_xml.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                        rel_id = rel.attrib.get("Id")
                        if not rel_id:
                            continue
                        rels[rel_id] = {
                            "type": rel.attrib.get("Type", ""),
                            "target": rel.attrib.get("Target", ""),
                            "mode": rel.attrib.get("TargetMode", ""),
                        }
                    part_rels[owner] = rels
            for name in zf.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    media[name.split("word/", 1)[1]] = zf.read(name)

        return document_xml, styles_xml, numbering_xml, part_rels, media

    def _parse_styles(self):
        styles = {}
        if self.styles_xml is None:
            return styles

        for style in self.styles_xml.findall("./w:style", NS):
            style_id = style.attrib.get(f"{{{W_NS}}}styleId")
            if not style_id:
                continue

            name_node = style.find("./w:name", NS)
            ppr = style.find("./w:pPr", NS)
            outline = None
            if ppr is not None:
                outline_node = ppr.find("./w:outlineLvl", NS)
                if outline_node is not None:
                    outline = outline_node.attrib.get(f"{{{W_NS}}}val")

            styles[style_id] = {
                "name": name_node.attrib.get(f"{{{W_NS}}}val") if name_node is not None else style_id,
                "outline": outline,
            }

        return styles

    def _parse_numbering(self):
        numbering = {"nums": {}, "abstract": {}}
        if self.numbering_xml is None:
            return numbering

        for abstract in self.numbering_xml.findall("./w:abstractNum", NS):
            abstract_id = abstract.attrib.get(f"{{{W_NS}}}abstractNumId")
            if abstract_id is None:
                continue

            levels = {}
            for lvl in abstract.findall("./w:lvl", NS):
                ilvl = lvl.attrib.get(f"{{{W_NS}}}ilvl")
                if ilvl is None:
                    continue

                num_fmt = lvl.find("./w:numFmt", NS)
                lvl_text = lvl.find("./w:lvlText", NS)
                start = lvl.find("./w:start", NS)
                levels[int(ilvl)] = {
                    "numFmt": num_fmt.attrib.get(f"{{{W_NS}}}val") if num_fmt is not None else "decimal",
                    "lvlText": lvl_text.attrib.get(f"{{{W_NS}}}val") if lvl_text is not None else "%1.",
                    "start": int(start.attrib.get(f"{{{W_NS}}}val", "1")) if start is not None else 1,
                }

            numbering["abstract"][int(abstract_id)] = levels

        for num in self.numbering_xml.findall("./w:num", NS):
            num_id = num.attrib.get(f"{{{W_NS}}}numId")
            abstract_num = num.find("./w:abstractNumId", NS)
            if num_id is None or abstract_num is None:
                continue
            numbering["nums"][int(num_id)] = int(abstract_num.attrib.get(f"{{{W_NS}}}val", "0"))

        return numbering

    def _resolve_media_html(self, rel_id: str, rels: dict[str, dict[str, str]]) -> str:
        rel = rels.get(rel_id)
        if not rel:
            return ""
        target = rel.get("target", "").lstrip("/")
        if not target:
            return ""
        media_name = target if target.startswith("word/") else f"word/{target}"
        media_name = media_name.split("word/", 1)[1]
        data = self.media.get(media_name)
        if not data:
            return ""
        mime = mimetypes.guess_type(media_name)[0] or "application/octet-stream"
        return f'<img src="data:{mime};base64,{base64.b64encode(data).decode("ascii")}" style="max-width:100%; height:auto;" />'

    def _escape_with_placeholders(self, value: str) -> str:
        if not value:
            return ""

        parts: list[str] = []
        index = 0
        for match in PLACEHOLDER_RE.finditer(value):
            parts.append(escape(value[index:match.start()]))
            parts.append(f"{{{{{_normalize_placeholder_name(match.group(1))}}}}}")
            index = match.end()
        parts.append(escape(value[index:]))
        return "".join(parts)

    def _paragraph_tab_stops(self, paragraph_elem) -> list[int]:
        ppr = paragraph_elem.find("./w:pPr", NS)
        if ppr is None:
            return []

        stops: list[int] = []
        tabs = ppr.find("./w:tabs", NS)
        if tabs is None:
            return []

        for tab in tabs.findall("./w:tab", NS):
            pos = tab.attrib.get(f"{{{W_NS}}}pos")
            if not pos:
                continue
            try:
                stops.append(max(int(pos) // 20, 0))
            except ValueError:
                continue
        return stops

    def _run_css(self, run_elem) -> str:
        rpr = run_elem.find("./w:rPr", NS)
        styles = []
        if rpr is None:
            return ""

        color = rpr.find("./w:color", NS)
        if color is not None:
            val = color.attrib.get(f"{{{W_NS}}}val")
            if val and val.lower() != "auto":
                styles.append(f"color:#{val}")

        size = rpr.find("./w:sz", NS)
        size_pt = _half_points_to_pt(size.attrib.get(f"{{{W_NS}}}val")) if size is not None else None
        if size_pt:
            styles.append(f"font-size:{size_pt:g}pt")

        rfonts = rpr.find("./w:rFonts", NS)
        if rfonts is not None:
            font = rfonts.attrib.get(f"{{{W_NS}}}ascii") or rfonts.attrib.get(f"{{{W_NS}}}hAnsi")
            if font:
                styles.append(f"font-family:'{font}'")

        return "; ".join(styles)

    def _render_run_segments(self, run_elem, rels: dict[str, dict[str, str]] | None = None) -> list[str]:
        rels = rels or self.rels
        rpr = run_elem.find("./w:rPr", NS)
        bold = rpr.find("./w:b", NS) is not None if rpr is not None else False
        italic = rpr.find("./w:i", NS) is not None if rpr is not None else False
        underline = rpr.find("./w:u", NS) is not None if rpr is not None else False

        text_parts: list[str] = []
        segments: list[str] = []
        image_html: Optional[str] = None

        def flush_segment() -> None:
            nonlocal text_parts, image_html
            text = "".join(text_parts)
            if text or image_html:
                rendered = self._escape_with_placeholders(text).replace("\t", "&emsp;").replace("\n", "<br/>")
                if underline:
                    rendered = f"<u>{rendered}</u>"
                if italic:
                    rendered = f"<em>{rendered}</em>"
                if bold:
                    rendered = f"<strong>{rendered}</strong>"

                css = self._run_css(run_elem)
                if css:
                    rendered = f'<span style="{css}">{rendered}</span>'
                if image_html:
                    rendered = f"{rendered}{image_html}"
                segments.append(rendered)
            text_parts = []
            image_html = None

        for child in list(run_elem):
            tag = _local_name(child)
            if tag == "t":
                text_parts.append(child.text or "")
            elif tag == "tab":
                flush_segment()
            elif tag in {"br", "cr"}:
                text_parts.append("\n")
            elif tag == "drawing":
                blip = child.find(".//a:blip", NS)
                if blip is not None:
                    rel_id = blip.attrib.get(f"{{{R_NS}}}embed")
                    if rel_id:
                        image_html = self._resolve_media_html(rel_id, rels)

        flush_segment()
        return segments

    def _render_paragraph_content(self, paragraph_elem, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        if self._paragraph_has_tabs(paragraph_elem):
            chunks = self._paragraph_tab_chunks(paragraph_elem)
            if len(chunks) > 1 and not self._should_row_layout(paragraph_elem, chunks):
                return "&emsp;".join(chunk for chunk in chunks if str(chunk).strip())

        parts: list[str] = []
        for child in list(paragraph_elem):
            tag = _local_name(child)
            if tag == "r":
                for rendered in self._render_run_segments(child, rels=rels):
                    if rendered:
                        parts.append(rendered)
            elif tag == "hyperlink":
                rel_id = child.attrib.get(f"{{{R_NS}}}id")
                href = rels.get(rel_id, {}).get("target") if rel_id else None
                link_parts = []
                for run in child.findall("./w:r", NS):
                    link_parts.extend(self._render_run_segments(run, rels=rels))
                if link_parts:
                    label = "".join(link_parts)
                    if href and rels.get(rel_id, {}).get("mode") == "External":
                        parts.append(f'<a href="{escape(href, quote=True)}">{label}</a>')
                    else:
                        parts.append(label)
            elif tag == "smartTag":
                for run in child.findall(".//w:r", NS):
                    parts.extend(self._render_run_segments(run, rels=rels))
        return "".join(parts)

    def _paragraph_css(self, paragraph_elem) -> str:
        ppr = paragraph_elem.find("./w:pPr", NS)
        styles = ["margin:0 0 10px 0", "white-space:pre-wrap"]
        if ppr is None:
            return "; ".join(styles)

        jc = ppr.find("./w:jc", NS)
        if jc is not None:
            align = jc.attrib.get(f"{{{W_NS}}}val") or jc.attrib.get("val")
            if align == "center":
                styles.append("text-align:center")
            elif align == "right":
                styles.append("text-align:right")
            elif align == "both":
                styles.append("text-align:justify")
            else:
                styles.append("text-align:left")

        spacing = ppr.find("./w:spacing", NS)
        if spacing is not None:
            before = spacing.attrib.get(f"{{{W_NS}}}before")
            after = spacing.attrib.get(f"{{{W_NS}}}after")
            line = spacing.attrib.get(f"{{{W_NS}}}line")
            if before:
                styles.append(f"margin-top:{_twips_to_px(before)}px")
            if after:
                styles.append(f"margin-bottom:{_twips_to_px(after)}px")
            if line:
                try:
                    styles.append(f"line-height:{max(int(line) / 240.0, 1.0):.2f}")
                except ValueError:
                    pass

        ind = ppr.find("./w:ind", NS)
        if ind is not None:
            left = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}left"))
            first = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}firstLine"))
            hanging = _twips_to_px(ind.attrib.get(f"{{{W_NS}}}hanging"))
            if left:
                styles.append(f"margin-left:{left}px")
            if first:
                styles.append(f"text-indent:{first}px")
            if hanging:
                styles.append(f"text-indent:-{hanging}px")
                if left:
                    styles.append(f"margin-left:{left + hanging}px")

        return "; ".join(styles)

    def _heading_level(self, paragraph_elem) -> int:
        ppr = paragraph_elem.find("./w:pPr", NS)
        pstyle = paragraph_elem.find("./w:pPr/w:pStyle", NS)
        style_id = pstyle.attrib.get(f"{{{W_NS}}}val") if pstyle is not None else None
        style = self.styles.get(style_id or "", {})
        name = (style.get("name") or "").lower()
        if name.startswith("heading "):
            try:
                return min(max(int(name.split("heading ", 1)[1]), 1), 6)
            except ValueError:
                return 0

        if style.get("outline") is not None:
            try:
                return min(int(style["outline"]) + 1, 6)
            except ValueError:
                return 0

        if pstyle is not None and style_id and style_id.lower().startswith("heading"):
            digits = re.sub(r"\D+", "", style_id)
            if digits:
                return min(max(int(digits), 1), 6)

        return 0

    def _list_format(self, num_id: int, ilvl: int) -> ListFormat:
        abstract_id = self.numbering["nums"].get(num_id)
        fmt = self.numbering["abstract"].get(abstract_id or -1, {}).get(ilvl, {})
        num_fmt = (fmt.get("numFmt") or "decimal").lower()
        if num_fmt == "bullet":
            return ListFormat("ul")
        if num_fmt in {"lowerletter", "alpha"}:
            return ListFormat("ol", ' type="a"')
        if num_fmt in {"upperletter"}:
            return ListFormat("ol", ' type="A"')
        if num_fmt in {"lowerroman"}:
            return ListFormat("ol", ' type="i"')
        if num_fmt in {"upperroman"}:
            return ListFormat("ol", ' type="I"')
        return ListFormat("ol", ' type="1"')

    def _paragraph_list_meta(self, paragraph_elem):
        num_pr = paragraph_elem.find("./w:pPr/w:numPr", NS)
        if num_pr is None:
            return None
        num_id = num_pr.find("./w:numId", NS)
        ilvl = num_pr.find("./w:ilvl", NS)
        if num_id is None or ilvl is None:
            return None
        try:
            return int(num_id.attrib.get(f"{{{W_NS}}}val", "0")), int(ilvl.attrib.get(f"{{{W_NS}}}val", "0"))
        except ValueError:
            return None

    def _render_table_cell(self, cell_elem, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        blocks: list[str] = []
        for child in list(cell_elem):
            tag = _local_name(child)
            if tag == "p":
                content = self._render_paragraph_content(child, rels=rels)
                if content.strip():
                    blocks.append(f'<p style="{self._paragraph_css(child)}">{content}</p>')
            elif tag == "tbl":
                blocks.append(self._render_table(child))
        return "".join(blocks) or "&nbsp;"

    def _table_has_visible_borders(self, table_elem) -> bool:
        tbl_pr = table_elem.find("./w:tblPr", NS)
        if tbl_pr is None:
            return False

        borders = tbl_pr.find("./w:tblBorders", NS)
        if borders is None:
            return False

        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = borders.find(f"./w:{side}", NS)
            if border is None:
                continue
            value = (border.attrib.get(f"{{{W_NS}}}val") or border.attrib.get("val") or "").lower()
            if value and value not in {"nil", "none", "0"}:
                return True
        return False

    def _paragraph_segments(self, paragraph_elem) -> list[str]:
        segments = [""]

        for child in list(paragraph_elem):
            tag = _local_name(child)
            if tag == "r":
                pieces = self._render_run_segments(child)
            elif tag == "hyperlink":
                pieces = []
                for run in child.findall("./w:r", NS):
                    pieces.extend(self._render_run_segments(run))
            elif tag == "smartTag":
                pieces = []
                for run in child.findall(".//w:r", NS):
                    pieces.extend(self._render_run_segments(run))
            else:
                continue

            if not pieces:
                continue

            segments[-1] += pieces[0]
            for piece in pieces[1:]:
                segments.append(piece)

        return [segment for segment in segments if segment is not None]

    def _paragraph_has_tabs(self, paragraph_elem) -> bool:
        if paragraph_elem.find(".//w:tab", NS) is not None:
            return True
        return bool(self._paragraph_tab_stops(paragraph_elem))

    def _paragraph_has_hanging_indent(self, paragraph_elem) -> bool:
        ppr = paragraph_elem.find("./w:pPr", NS)
        if ppr is None:
            return False
        ind = ppr.find("./w:ind", NS)
        if ind is None:
            return False
        hanging = ind.attrib.get(f"{{{W_NS}}}hanging")
        return bool(hanging)

    def _paragraph_tab_chunks(self, paragraph_elem) -> list[str]:
        chunks = [""]

        def append_piece(piece: str) -> None:
            nonlocal chunks
            if piece == "\t":
                chunks.append("")
                return
            if piece:
                chunks[-1] += piece

        for child in list(paragraph_elem):
            tag = _local_name(child)
            if tag == "r":
                for inner in list(child):
                    inner_tag = _local_name(inner)
                    if inner_tag == "t":
                        append_piece(self._escape_with_placeholders(inner.text or ""))
                    elif inner_tag == "tab":
                        append_piece("\t")
                    elif inner_tag in {"br", "cr"}:
                        append_piece("<br/>")
                    elif inner_tag == "drawing":
                        blip = inner.find(".//a:blip", NS)
                        if blip is not None:
                            rel_id = blip.attrib.get(f"{{{R_NS}}}embed")
                            if rel_id and rel_id in self.rels:
                                target = self.rels[rel_id]["target"].lstrip("/")
                                media_name = target if target.startswith("word/") else f"word/{target}"
                                media_name = media_name.split("word/", 1)[1]
                                data = self.media.get(media_name)
                                if data:
                                    mime = mimetypes.guess_type(media_name)[0] or "application/octet-stream"
                                    append_piece(f'<img src="data:{mime};base64,{base64.b64encode(data).decode("ascii")}" style="max-width:100%; height:auto;" />')
            elif tag in {"hyperlink", "smartTag"}:
                for run in child.findall(".//w:r", NS):
                    for inner in list(run):
                        inner_tag = _local_name(inner)
                        if inner_tag == "t":
                            append_piece(self._escape_with_placeholders(inner.text or ""))
                        elif inner_tag == "tab":
                            append_piece("\t")
                        elif inner_tag in {"br", "cr"}:
                            append_piece("<br/>")

        return [chunk for chunk in chunks if chunk is not None]

    def _is_label_token(self, value: str) -> bool:
        text = re.sub(r"\s+", " ", value).strip()
        if not text:
            return False
        if len(text) > 40:
            return False
        if text.endswith(":"):
            return True
        return bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 ,./&()\-]{0,38}", text))

    def _is_value_token(self, value: str) -> bool:
        text = re.sub(r"\s+", " ", value).strip()
        if not text:
            return False
        if re.fullmatch(r"[\d\s\-/.,()]+", text):
            return True
        if re.match(r"^date\s*:\s*", text, flags=re.IGNORECASE):
            return True
        if len(text) <= 30 and "." not in text and " " not in text:
            return True
        return False

    def _should_row_layout(self, paragraph_elem, chunks: list[str]) -> bool:
        chunks = [chunk for chunk in chunks if str(chunk).strip()]
        if len(chunks) < 2:
            return False

        first_visible = re.sub(r"\s+", " ", str(chunks[0]).strip())
        if re.match(r"^(?:[0-9]+|[a-dA-D])[\.\)]\s*$", first_visible):
            return False

        hanging_indent = self._paragraph_has_hanging_indent(paragraph_elem)
        first_is_label = self._is_label_token(first_visible)
        last_visible = re.sub(r"\s+", " ", str(chunks[-1]).strip())
        last_is_value = self._is_value_token(last_visible)

        if hanging_indent:
            return first_is_label and last_is_value and len(chunks) == 2

        if len(chunks) == 2:
            return first_is_label and last_is_value

        # Multi-chunk tabbed paragraphs are usually structured rows when they
        # do not look like manual list numbering.
        return first_is_label

    def _row_layout_from_chunks(self, paragraph_elem, chunks: list[str]) -> str:
        chunks = [chunk for chunk in chunks if str(chunk).strip()]
        if not self._should_row_layout(paragraph_elem, chunks):
            return ""

        tab_stops = self._paragraph_tab_stops(paragraph_elem)
        widths: list[str] = []
        if tab_stops and len(tab_stops) >= len(chunks) - 1:
            for i, _ in enumerate(chunks):
                if i == 0:
                    widths.append(f"{tab_stops[0]}px")
                elif i < len(chunks) - 1:
                    widths.append(f"{max(tab_stops[i] - tab_stops[i - 1], 1)}px")
                else:
                    widths.append("auto")
        else:
            widths = ["auto"] * len(chunks)

        cells = []
        for idx, chunk in enumerate(chunks):
            text = chunk.replace("\t", "")
            align = "left"
            if idx == len(chunks) - 1 and len(chunks) > 1:
                align = "right" if self._is_value_token(text) else "left"
            width_css = f"width:{widths[idx]};" if idx < len(widths) and widths[idx] != "auto" else ""
            cells.append(
                f'<td class="{"right" if align == "right" else "left"}" style="{width_css} text-align:{align}; vertical-align:top; padding:0;">{text or "&nbsp;"}</td>'
            )

        return f'<table class="row-layout" style="width:100%; border-collapse:collapse; table-layout:fixed;"><tr>{"".join(cells)}</tr></table>'

    def _render_row_layout(self, paragraph_elem) -> str:
        chunks = self._paragraph_tab_chunks(paragraph_elem)
        return self._row_layout_from_chunks(paragraph_elem, chunks)

    def _render_table(self, table_elem) -> str:
        tbl_pr = table_elem.find("./w:tblPr", NS)
        tbl_grid = table_elem.find("./w:tblGrid", NS)
        visible_borders = self._table_has_visible_borders(table_elem)
        col_widths = []
        if tbl_grid is not None:
            for col in tbl_grid.findall("./w:gridCol", NS):
                width = col.attrib.get(f"{{{W_NS}}}w")
                if width:
                    col_widths.append(max(int(width) // 20, 1))

        rows_html = []
        tr_elems = [child for child in list(table_elem) if _local_name(child) == "tr"]
        for row_idx, row in enumerate(tr_elems):
            row_cells = []
            tc_elems = [child for child in list(row) if _local_name(child) == "tc"]
            is_header_row = row_idx == 0 or row.find("./w:trPr/w:tblHeader", NS) is not None
            for cell_idx, cell in enumerate(tc_elems):
                tc_pr = cell.find("./w:tcPr", NS)
                colspan = 1
                rowspan = None
                if tc_pr is not None:
                    grid_span = tc_pr.find("./w:gridSpan", NS)
                    if grid_span is not None:
                        try:
                            colspan = max(int(grid_span.attrib.get(f"{{{W_NS}}}val", "1")), 1)
                        except ValueError:
                            colspan = 1
                    vmerge = tc_pr.find("./w:vMerge", NS)
                    if vmerge is not None:
                        vmerge_val = vmerge.attrib.get(f"{{{W_NS}}}val")
                        if vmerge_val == "restart" or vmerge_val is None:
                            rowspan = 1
                        else:
                            continue

                content = self._render_table_cell(cell, rels=self.rels)
                tag = "th" if is_header_row else "td"
                align = "center" if is_header_row else ("right" if cell_idx == len(tc_elems) - 1 else "left")
                if tc_pr is not None:
                    jc = tc_pr.find("./w:jc", NS)
                    if jc is not None:
                        align = jc.attrib.get(f"{{{W_NS}}}val") or align

                cell_width_style = ""
                if tc_pr is not None:
                    tc_w = tc_pr.find("./w:tcW", NS)
                    if tc_w is not None:
                        width_val = tc_w.attrib.get(f"{{{W_NS}}}w")
                        width_type = tc_w.attrib.get(f"{{{W_NS}}}type")
                        if width_val:
                            if width_type == "pct":
                                try:
                                    cell_width_style = f"width:{int(width_val) / 50:.2f}%"
                                except ValueError:
                                    cell_width_style = ""
                            else:
                                try:
                                    cell_width_style = f"width:{max(int(width_val) // 20, 1)}px"
                                except ValueError:
                                    cell_width_style = ""

                border_style = "border:1px solid #111827;" if visible_borders else "border:none;"
                attrs = [
                    'style="%s padding:8px 10px; vertical-align:top; text-align:%s; line-height:1.2;%s"'
                    % (border_style, align, f" {cell_width_style};" if cell_width_style else "")
                ]
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')
                if rowspan:
                    attrs.append(f'rowspan="{rowspan}"')
                row_cells.append(f"<{tag} {' '.join(attrs)}>{content}</{tag}>")
            rows_html.append(f"<tr>{''.join(row_cells)}</tr>")

        table_style = "width:100%; border-collapse:collapse; table-layout:fixed; margin:10px 0; page-break-inside:avoid;"
        colgroup = ""
        if col_widths:
            cols = "".join(f'<col style="width:{w}px" />' for w in col_widths)
            colgroup = f"<colgroup>{cols}</colgroup>"

        table_class = "docx-table" if visible_borders else "docx-table hidden-table"
        return f'<table class="{table_class}" style="{table_style}">{colgroup}{"".join(rows_html)}</table>'

    def _length_to_px(self, value) -> int:
        if value is None:
            return 0
        try:
            if hasattr(value, "pt") and value.pt is not None:
                return max(int(round(float(value.pt) * 1.333333)), 0)
        except Exception:
            pass
        try:
            return max(int(value), 0)
        except Exception:
            return 0

    def _docx_alignment_css(self, alignment) -> str:
        if alignment in (WD_ALIGN_PARAGRAPH.CENTER,):
            return "text-align:center"
        if alignment in (WD_ALIGN_PARAGRAPH.RIGHT,):
            return "text-align:right"
        if alignment in (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.DISTRIBUTE):
            return "text-align:justify"
        return "text-align:left"

    def _docx_run_html(self, run, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        text = run.text or ""
        if not text:
            text = ""

        rendered = self._escape_with_placeholders(text).replace("\t", "&emsp;").replace("\n", "<br/>")
        image_html = ""
        for child in list(run.element):
            tag = _local_name(child)
            if tag == "drawing":
                blip = child.find(".//a:blip", NS)
                if blip is not None:
                    rel_id = blip.attrib.get(f"{{{R_NS}}}embed")
                    if rel_id:
                        image_html = self._resolve_media_html(rel_id, rels)
        if run.underline:
            rendered = f"<u>{rendered}</u>"
        if run.italic:
            rendered = f"<em>{rendered}</em>"
        if run.bold:
            rendered = f"<strong>{rendered}</strong>"

        styles = []
        if getattr(run.font, "name", None):
            styles.append(f"font-family:'{run.font.name}'")
        if getattr(run.font, "size", None) is not None:
            size_px = self._length_to_px(run.font.size)
            if size_px:
                styles.append(f"font-size:{size_px}px")
        if styles:
            rendered = f'<span style="{"; ".join(styles)}">{rendered}</span>'
        if image_html:
            rendered = f"{rendered}{image_html}"
        return rendered

    def _docx_paragraph_html(self, paragraph, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        parts: list[str] = []
        for run in paragraph.runs:
            run_html = self._docx_run_html(run, rels=rels)
            if run_html:
                parts.append(run_html)
        content = "".join(parts)
        if not content.strip():
            return ""

        style_bits = ["margin:0 0 10px 0", "white-space:pre-wrap"]
        style_bits.append(self._docx_alignment_css(paragraph.alignment))

        fmt = paragraph.paragraph_format
        if fmt is not None:
            if fmt.left_indent is not None:
                style_bits.append(f"margin-left:{self._length_to_px(fmt.left_indent)}px")
            if fmt.first_line_indent is not None:
                indent_px = self._length_to_px(fmt.first_line_indent)
                style_bits.append(f"text-indent:{indent_px}px")
            if fmt.space_before is not None:
                style_bits.append(f"margin-top:{self._length_to_px(fmt.space_before)}px")
            if fmt.space_after is not None:
                style_bits.append(f"margin-bottom:{self._length_to_px(fmt.space_after)}px")
            if fmt.line_spacing is not None:
                try:
                    style_bits.append(f"line-height:{float(fmt.line_spacing):.2f}")
                except Exception:
                    pass

        return f'<p style="{"; ".join(style_bits)}">{content}</p>'

    def _docx_table_html(self, table, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        rows = []
        colgroup = ""
        if table.columns:
            widths = []
            for column in table.columns:
                width = getattr(column, "width", None)
                px = self._length_to_px(width)
                widths.append(px if px else None)
            if any(widths):
                cols = "".join(
                    f'<col style="width:{width}px" />' if width else "<col />"
                    for width in widths
                )
                colgroup = f"<colgroup>{cols}</colgroup>"

        for row_idx, row in enumerate(table.rows):
            cells = []
            for cell_idx, cell in enumerate(row.cells):
                cell_parts = []
                for paragraph in cell.paragraphs:
                    rendered = self._docx_paragraph_html(paragraph, rels=rels)
                    if rendered:
                        cell_parts.append(rendered)
                cell_html = "".join(cell_parts) or "&nbsp;"
                cell_align = "right" if cell_idx == len(row.cells) - 1 else "left"
                cells.append(
                    f'<td style="border:1px solid #111827; padding:8px 10px; vertical-align:top; text-align:{cell_align}; line-height:1.2;">{cell_html}</td>'
                )
            rows.append(f"<tr>{''.join(cells)}</tr>")

        return f'<table class="docx-table" style="width:100%; border-collapse:collapse; table-layout:fixed; margin:10px 0; page-break-inside:avoid;">{colgroup}{"".join(rows)}</table>'

    def _render_docx_container(self, container, rels: dict[str, dict[str, str]] | None = None) -> str:
        rels = rels or self.rels
        blocks: list[str] = []
        for paragraph in getattr(container, "paragraphs", []):
            rendered = self._docx_paragraph_html(paragraph, rels=rels)
            if rendered:
                blocks.append(rendered)
        for table in getattr(container, "tables", []):
            rendered = self._docx_table_html(table, rels=rels)
            if rendered:
                blocks.append(rendered)
        return "\n".join(blocks)

    def _render_header_footer_blocks(self) -> tuple[str, str]:
        header_blocks: list[str] = []
        footer_blocks: list[str] = []
        seen_headers = set()
        seen_footers = set()

        for section in self._doc.sections:
            header = section.header
            footer = section.footer
            header_rels = self.part_rels.get("word/" + section.header.part.partname.split("/")[-1], self.rels)
            footer_rels = self.part_rels.get("word/" + section.footer.part.partname.split("/")[-1], self.rels)

            header_key = id(getattr(header, "_element", header))
            footer_key = id(getattr(footer, "_element", footer))

            if header and header_key not in seen_headers:
                rendered_header = self._render_docx_container(header, rels=header_rels)
                if rendered_header.strip():
                    header_blocks.append(f'<header class="docx-header">{rendered_header}</header>')
                seen_headers.add(header_key)

            if footer and footer_key not in seen_footers:
                rendered_footer = self._render_docx_container(footer, rels=footer_rels)
                if rendered_footer.strip():
                    footer_blocks.append(f'<footer class="docx-footer">{rendered_footer}</footer>')
                seen_footers.add(footer_key)

        return "\n".join(header_blocks), "\n".join(footer_blocks)

    def _render_block_paragraph(self, paragraph_elem) -> str:
        content = self._render_paragraph_content(paragraph_elem)
        if not content.strip():
            return ""
        level = self._heading_level(paragraph_elem)
        css = self._paragraph_css(paragraph_elem)
        if level:
            return f"<h{level} style=\"{css}\">{content}</h{level}>"
        return f'<p style="{css}">{content}</p>'

    def render(self) -> str:
        body = self.document_xml.find("./w:body", NS)
        if body is None:
            return ""

        parts: list[str] = []
        list_stack: list[dict[str, object]] = []

        def close_all_lists():
            while list_stack:
                ctx = list_stack.pop()
                if ctx["li_open"]:
                    parts.append("</li>")
                parts.append(f'</{ctx["tag"]}>')

        for child in list(body):
            tag = _local_name(child)
            if tag == "p":
                list_meta = self._paragraph_list_meta(child)
                if list_meta:
                    num_id, ilvl = list_meta
                    level = ilvl + 1
                    fmt = self._list_format(num_id, ilvl)

                    while len(list_stack) > level:
                        ctx = list_stack.pop()
                        if ctx["li_open"]:
                            parts.append("</li>")
                        parts.append(f'</{ctx["tag"]}>')

                    while len(list_stack) < level:
                        parts.append(f"<{fmt.tag}{fmt.type_attr}>")
                        list_stack.append({"tag": fmt.tag, "type": fmt.type_attr, "li_open": False})

                    ctx = list_stack[-1]
                    if ctx["li_open"]:
                        parts.append("</li>")
                        ctx["li_open"] = False

                    item_css = self._paragraph_css(child)
                    content = self._render_paragraph_content(child)
                    parts.append(f'<li style="{item_css}">{content}')
                    ctx["li_open"] = True
                    continue

                if self._paragraph_has_tabs(child):
                    chunks = self._paragraph_tab_chunks(child)
                    if self._should_row_layout(child, chunks):
                        close_all_lists()
                        row_layout = self._row_layout_from_chunks(child, chunks)
                        if row_layout:
                            parts.append(row_layout)
                            continue

                close_all_lists()
                rendered = self._render_block_paragraph(child)
                if rendered:
                    parts.append(rendered)
                continue

            if tag == "tbl":
                close_all_lists()
                parts.append(self._render_table(child))
                continue

            if tag == "sectPr":
                close_all_lists()
                parts.append('<div style="break-after:page;"></div>')

        close_all_lists()
        return "\n".join(parts)


def docx_to_html(file_path: str) -> str:
    renderer = DocxHtmlRenderer(file_path)
    body_html = renderer.render()
    header_html, footer_html = renderer._render_header_footer_blocks()
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <style>
    @page {{
      size: A4;
      margin: 20mm;
    }}

    * {{
      box-sizing: border-box;
    }}

    body {{
      margin: 0;
      font-family: Arial, Helvetica, sans-serif;
      font-size: 12pt;
      line-height: 1.5;
      color: #111827;
      background: #fff;
    }}

    .docx-page {{
      width: 100%;
    }}

    .docx-header,
    .docx-footer {{
      width: 100%;
      overflow: hidden;
    }}

    .docx-header {{
      margin: 0 0 14px 0;
    }}

    .docx-footer {{
      margin: 14px 0 0 0;
    }}

    .docx-table {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
    }}

    .docx-table.hidden-table,
    .docx-table.hidden-table th,
    .docx-table.hidden-table td {{
      border: none;
    }}

    .row-layout {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
    }}

    .row-layout .left {{
      text-align: left;
    }}

    .row-layout .right {{
      text-align: right;
    }}

    .table-fixed {{
      width: 100%;
      border-collapse: collapse;
    }}

    .docx-table th,
    .docx-table td {{
      border: 1px solid #111827;
      vertical-align: top;
    }}

    p, li, h1, h2, h3, h4, h5, h6 {{
      margin: 0;
      white-space: pre-wrap;
    }}

    ol, ul {{
      margin: 0;
      padding-left: 24px;
    }}

    img {{
      max-width: 100%;
      height: auto;
    }}

    .docx-header img {{
      max-height: 24mm;
      width: auto;
    }}

    .docx-footer img {{
      max-height: 18mm;
      width: auto;
    }}

    @media print {{
      @page {{
        size: A4;
        margin: 36mm 18mm 30mm 18mm;
      }}

      body {{
        margin: 0;
        padding-top: 36mm;
        padding-bottom: 30mm;
      }}

      .docx-header {{
        position: fixed;
        top: 0;
        left: 18mm;
        right: 18mm;
        height: 28mm;
        margin: 0;
        padding: 0;
        overflow: hidden;
      }}

      .docx-footer {{
        position: fixed;
        bottom: 0;
        left: 18mm;
        right: 18mm;
        height: 22mm;
        margin: 0;
        padding: 0;
        overflow: hidden;
      }}

      .docx-page {{
        padding-top: 0;
        padding-bottom: 0;
      }}

      .docx-header img {{
        max-height: 26mm;
      }}

      .docx-footer img {{
        max-height: 20mm;
      }}
    }}
  </style>
</head>
<body>
  <div class="docx-page">
    {header_html}
    {body_html}
    {footer_html}
  </div>
</body>
</html>"""
